"""Text out of an attachment, for the grading prompt.

Images produce no text at all: they are not transcribed, they travel to the model as
images. Everything else becomes text, capped — one spreadsheet can be larger than every
other part of the prompt combined, and a prompt that overruns the window does not error,
it silently drops the overflow.

Caps are per file and in total. Both are reported to the model when they bite, so it can
tell "the answer is short" from "we only sent the first 6000 characters".
"""

import io

from app.attachments.storage import AttachmentError

MAX_CHARS_PER_FILE = 6000
MAX_CHARS_TOTAL = 20000
MAX_PDF_PAGES = 20
MAX_SHEET_ROWS = 500


def extract_text(kind: str, data: bytes) -> str:
    """Plain text for one attachment; empty for images."""
    if kind == "pdf":
        return _pdf(data)
    if kind == "docx":
        return _docx(data)
    if kind == "xlsx":
        return _xlsx(data)
    if kind == "text":
        return _text(data)
    return ""


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001 — pypdf raises several unrelated types
        raise AttachmentError("这个 PDF 读不出来，可能已损坏。") from exc
    pages: list[str] = []
    for page in reader.pages[:MAX_PDF_PAGES]:
        pages.append(page.extract_text() or "")
    return "\n\n".join(pages)


def _docx(data: bytes) -> str:
    # Reuses the resource parser's hand-rolled DOCX reader, so an attachment and a
    # course resource go through the same extraction rather than two that disagree.
    from app.resources.parser import parse_docx_chunks

    try:
        chunks = parse_docx_chunks(data)
    except ValueError as exc:
        raise AttachmentError("这个 Word 文件读不出来，可能已损坏。") from exc
    return "\n\n".join(chunk.text for chunk in chunks)


def _xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 — openpyxl raises several unrelated types
        raise AttachmentError("这个 Excel 文件读不出来，可能已损坏。") from exc
    lines: list[str] = []
    try:
        for sheet in book.worksheets:
            lines.append(f"[sheet: {sheet.title}]")
            for index, row in enumerate(sheet.iter_rows(values_only=True)):
                if index >= MAX_SHEET_ROWS:
                    lines.append(f"…（本表仅读取前 {MAX_SHEET_ROWS} 行）")
                    break
                cells = ["" if cell is None else str(cell) for cell in row]
                if any(cells):
                    lines.append("\t".join(cells))
    finally:
        book.close()
    return "\n".join(lines)


def _text(data: bytes) -> str:
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError as exc:
        raise AttachmentError("这个文本文件不是 UTF-8 编码，请另存为 UTF-8。") from exc


def cap(text: str, limit: int = MAX_CHARS_PER_FILE) -> str:
    """Truncate, and say so, rather than handing the model a silently cut document."""
    if len(text) <= limit:
        return text
    return f"{text[:limit]}\n…（已截断，原文共 {len(text)} 字）"


def describe(rows: list[tuple[str, str, str]]) -> str:
    """One prompt block from (filename, kind, text-or-empty) triples.

    Images are named but not inlined: they are attached to the request as images, and
    saying so here is what lets the model tell "the answer is in the picture" from
    "the learner attached a picture and wrote nothing".
    """
    if not rows:
        return ""
    blocks: list[str] = []
    remaining = MAX_CHARS_TOTAL
    for index, (filename, kind, text) in enumerate(rows, start=1):
        if kind == "image":
            blocks.append(f"[{index}] {filename}（图片，内容见随附图像）")
            continue
        if remaining <= 0:
            blocks.append(f"[{index}] {filename}（附件总长度已达上限，未发送内容）")
            continue
        body = cap(text, min(MAX_CHARS_PER_FILE, remaining))
        remaining -= len(body)
        blocks.append(f"[{index}] {filename}\n{body}")
    return "\n\n".join(blocks)
